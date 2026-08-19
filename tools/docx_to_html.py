"""
Convert a .docx into a standalone HTML file that visually matches it
(fonts, sizes, bold/italic/underline, alignment, spacing, indents, page
margins), while keeping any {{ jinja }} / {% jinja %} placeholders intact
as contiguous text so they can be filled in downstream (e.g. by n8n).

Usage:
    python docx_to_html.py <input.docx> [output.html]
"""
import os
import sys

import docx
from docx.oxml.ns import qn
import html as htmlmod

if len(sys.argv) < 2:
    print(__doc__)
    sys.exit(1)

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(SRC)[0] + '.html'

d = docx.Document(SRC)

try:
    doc_title = d.core_properties.title or os.path.splitext(os.path.basename(SRC))[0]
except Exception:
    doc_title = os.path.splitext(os.path.basename(SRC))[0]

# ---------- theme fonts (fallback defaults) ----------
theme_part = None
for rel in d.part.rels.values():
    if 'theme' in rel.reltype:
        theme_part = rel.target_part
        break
minor_latin = 'Calibri'
minor_eastasia = None
if theme_part is not None:
    from lxml import etree
    troot = etree.fromstring(theme_part.blob)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    minf = troot.find('.//a:fontScheme/a:minorFont/a:latin', ns)
    if minf is not None and minf.get('typeface'):
        minor_latin = minf.get('typeface')

DOC_DEFAULT = {
    'ascii': minor_latin, 'eastAsia': minor_eastasia, 'hAnsi': minor_latin,
    'bold': False, 'italic': False, 'underline': None, 'size_pt': 12.0, 'color': None,
    'jc': 'left', 'sp_before_pt': 0.0, 'sp_after_pt': 0.0,
    'line': None, 'lineRule': None, 'ind_left_pt': 0.0, 'ind_right_pt': 0.0,
}

def read_rstyle(rpr_el):
    if rpr_el is None:
        return None
    rstyle = rpr_el.find(qn('w:rStyle'))
    if rstyle is not None:
        return rstyle.get(qn('w:val'))
    return None

def read_rpr(rpr_el):
    """Extract only explicitly-set properties from a w:rPr element (excluding rStyle ref)."""
    out = {}
    if rpr_el is None:
        return out
    rfonts = rpr_el.find(qn('w:rFonts'))
    if rfonts is not None:
        if rfonts.get(qn('w:ascii')):
            out['ascii'] = rfonts.get(qn('w:ascii'))
        if rfonts.get(qn('w:eastAsia')):
            out['eastAsia'] = rfonts.get(qn('w:eastAsia'))
        if rfonts.get(qn('w:hAnsi')):
            out['hAnsi'] = rfonts.get(qn('w:hAnsi'))
    b = rpr_el.find(qn('w:b'))
    if b is not None:
        val = b.get(qn('w:val'))
        out['bold'] = False if val in ('0', 'false', 'off') else True
    i = rpr_el.find(qn('w:i'))
    if i is not None:
        val = i.get(qn('w:val'))
        out['italic'] = False if val in ('0', 'false', 'off') else True
    u = rpr_el.find(qn('w:u'))
    if u is not None:
        val = u.get(qn('w:val'))
        out['underline'] = None if val in (None, 'none') else val
    sz = rpr_el.find(qn('w:sz'))
    if sz is not None and sz.get(qn('w:val')):
        out['size_pt'] = int(sz.get(qn('w:val'))) / 2.0
    color = rpr_el.find(qn('w:color'))
    if color is not None:
        val = color.get(qn('w:val'))
        # '000000'/'auto' are indistinguishable from the default text color —
        # recording them as an "explicit" color would needlessly fragment
        # otherwise-identical runs (e.g. split a {{ placeholder }} in two).
        if val and val.lower() not in ('auto', '000000'):
            out['color'] = '#' + val
    return out

def read_ppr(ppr_el):
    out = {}
    if ppr_el is None:
        return out
    jc = ppr_el.find(qn('w:jc'))
    if jc is not None and jc.get(qn('w:val')):
        v = jc.get(qn('w:val'))
        out['jc'] = {'both': 'justify', 'start': 'left', 'end': 'right'}.get(v, v)
    spacing = ppr_el.find(qn('w:spacing'))
    if spacing is not None:
        if spacing.get(qn('w:before')) is not None:
            out['sp_before_pt'] = int(spacing.get(qn('w:before'))) / 20.0
        if spacing.get(qn('w:after')) is not None:
            out['sp_after_pt'] = int(spacing.get(qn('w:after'))) / 20.0
        if spacing.get(qn('w:line')) is not None:
            out['line'] = int(spacing.get(qn('w:line')))
            out['lineRule'] = spacing.get(qn('w:lineRule')) or 'auto'
    ind = ppr_el.find(qn('w:ind'))
    if ind is not None:
        if ind.get(qn('w:left')) is not None:
            out['ind_left_pt'] = int(ind.get(qn('w:left'))) / 20.0
        if ind.get(qn('w:right')) is not None:
            out['ind_right_pt'] = int(ind.get(qn('w:right'))) / 20.0
    return out

# ---------- resolve style chain (basedOn) ----------
styles_by_id = {}
for s in d.styles:
    sid = getattr(s, 'style_id', None)
    if sid:
        styles_by_id[sid] = s

_cache = {}
def resolve_style(style_id):
    if style_id in _cache:
        return _cache[style_id]
    if style_id is None or style_id not in styles_by_id:
        return dict(DOC_DEFAULT)
    style = styles_by_id[style_id]
    el = style.element
    based = el.find(qn('w:basedOn'))
    if based is not None:
        parent = resolve_style(based.get(qn('w:val')))
    else:
        parent = dict(DOC_DEFAULT)
    merged = dict(parent)
    merged.update(read_ppr(el.find(qn('w:pPr'))))
    merged.update(read_rpr(el.find(qn('w:rPr'))))
    _cache[style_id] = merged
    return merged

FONT_MAP = {
    '新細明體': 'PMingLiU',  # 新細明體
}
def css_font_family(effective):
    # Only `ascii` (Latin) / `hAnsi` govern rendering for Latin/placeholder
    # text; `eastAsia` is intentionally excluded here because when it's the
    # only difference between adjacent runs it would otherwise fragment a
    # {{ placeholder }} across multiple <span>s for no visible benefit.
    # If you convert a template with real CJK body text, add 'eastAsia'
    # back into this key list.
    fonts = []
    for key in ('ascii', 'hAnsi'):
        f = effective.get(key)
        if f:
            f = FONT_MAP.get(f, f)
            if f not in fonts:
                fonts.append(f)
    if not fonts:
        fonts = ['Calibri']
    parts = [f"'{f}'" if ' ' in f else f for f in fonts]
    parts.append('sans-serif')
    return ', '.join(parts)

def rpr_css(effective, base):
    """Return CSS declarations for properties in `effective` that differ from `base`."""
    decls = []
    fam_e, fam_b = css_font_family(effective), css_font_family(base)
    if fam_e != fam_b:
        decls.append(f'font-family: {fam_e}')
    if effective.get('size_pt') != base.get('size_pt'):
        decls.append(f'font-size: {effective.get("size_pt")}pt')
    if bool(effective.get('bold')) != bool(base.get('bold')):
        decls.append(f'font-weight: {"bold" if effective.get("bold") else "normal"}')
    if bool(effective.get('italic')) != bool(base.get('italic')):
        decls.append(f'font-style: {"italic" if effective.get("italic") else "normal"}')
    u_e, u_b = effective.get('underline'), base.get('underline')
    if u_e != u_b:
        decls.append(f'text-decoration: {"underline" if u_e else "none"}')
    c_e, c_b = effective.get('color'), base.get('color')
    if c_e != c_b:
        decls.append(f'color: {c_e or "inherit"}')
    return '; '.join(decls)

def p_css(effective):
    decls = []
    fam = css_font_family(effective)
    decls.append(f'font-family: {fam}')
    decls.append(f'font-size: {effective.get("size_pt")}pt')
    decls.append(f'font-weight: {"bold" if effective.get("bold") else "normal"}')
    if effective.get('italic'):
        decls.append('font-style: italic')
    if effective.get('underline'):
        decls.append('text-decoration: underline')
    if effective.get('color'):
        decls.append(f'color: {effective.get("color")}')
    decls.append(f'text-align: {effective.get("jc", "left")}')
    decls.append(f'margin: {effective.get("sp_before_pt", 0)}pt 0 {effective.get("sp_after_pt", 0)}pt 0')
    line, rule = effective.get('line'), effective.get('lineRule')
    if line:
        if rule == 'auto':
            decls.append(f'line-height: {round(line / 240.0, 3)}')
        else:
            decls.append(f'line-height: {round(line / 20.0, 2)}pt')
    else:
        decls.append('line-height: normal')
    if effective.get('ind_left_pt'):
        decls.append(f'margin-left: {effective.get("ind_left_pt")}pt')
    if effective.get('ind_right_pt'):
        decls.append(f'margin-right: {effective.get("ind_right_pt")}pt')
    return '; '.join(decls)

def esc(t):
    return htmlmod.escape(t, quote=False).replace(' ', '&nbsp;')

# ---------- page setup ----------
sec = d.sections[0]
page = {
    'w_in': sec.page_width.inches,
    'h_in': sec.page_height.inches,
    'top_in': sec.top_margin.inches,
    'bottom_in': sec.bottom_margin.inches,
    'left_in': sec.left_margin.inches,
    'right_in': sec.right_margin.inches,
}

body_parts = []
for p in d.paragraphs:
    style_id = p.style.style_id if p.style else None
    style_eff = resolve_style(style_id)
    direct_ppr = read_ppr(p._p.find(qn('w:pPr')))
    para_eff = dict(style_eff)
    para_eff.update(direct_ppr)

    segs = []  # list of (text, effective_rpr_dict, href_or_None)
    for item in p.iter_inner_content():
        cls_name = type(item).__name__
        if cls_name == 'Hyperlink':
            href = item.address
            for r in item.runs:
                rpr_el = r._r.find(qn('w:rPr'))
                rstyle_id = read_rstyle(rpr_el)
                r_eff = dict(style_eff)
                if rstyle_id:
                    r_eff.update(resolve_style(rstyle_id))
                r_eff.update(read_rpr(rpr_el))
                if r.text:
                    segs.append((r.text, r_eff, href))
        else:
            r = item
            rpr_el = r._r.find(qn('w:rPr'))
            rstyle_id = read_rstyle(rpr_el)
            r_eff = dict(style_eff)
            if rstyle_id:
                r_eff.update(resolve_style(rstyle_id))
            r_eff.update(read_rpr(rpr_el))
            if r.text:
                segs.append((r.text, r_eff, None))

    # merge adjacent segments whose *rendered* formatting (not every raw XML
    # attribute) is identical, so runs split only by irrelevant properties
    # (e.g. an eastAsia font that never applies to this all-Latin text) don't
    # fragment a {{ placeholder }} across multiple spans.
    def render_key(eff):
        return (css_font_family(eff), eff.get('size_pt'), bool(eff.get('bold')),
                bool(eff.get('italic')), eff.get('underline'), eff.get('color'))

    merged = []
    for text, eff, href in segs:
        key = render_key(eff)
        if merged and merged[-1][3] == key and merged[-1][2] == href:
            merged[-1] = (merged[-1][0] + text, merged[-1][1], href, key)
        else:
            merged.append((text, eff, href, key))
    merged = [(t, e, h) for t, e, h, k in merged]

    inner_html = []
    for text, eff, href in merged:
        css = rpr_css(eff, para_eff)
        content = esc(text)
        if href:
            style_attr = f' style="{css}"' if css else ''
            inner_html.append(f'<a href="{htmlmod.escape(href, quote=True)}"{style_attr}>{content}</a>')
        elif css:
            inner_html.append(f'<span style="{css}">{content}</span>')
        else:
            inner_html.append(content)

    html_content = ''.join(inner_html) if inner_html else '&nbsp;'
    p_style = p_css(para_eff)
    body_parts.append(f'<p style="{p_style}">{html_content}</p>')

css_page = f"""
@page {{
  size: {page['w_in']:.4f}in {page['h_in']:.4f}in;
  margin: {page['top_in']:.4f}in {page['right_in']:.4f}in {page['bottom_in']:.4f}in {page['left_in']:.4f}in;
}}
"""

html_doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>{htmlmod.escape(doc_title)}</title>
<style>
{css_page}
html, body {{
  margin: 0;
  padding: 0;
  background: #e6e6e6;
}}
.page {{
  box-sizing: border-box;
  width: {page['w_in']:.4f}in;
  min-height: {page['h_in']:.4f}in;
  margin: 24px auto;
  padding: {page['top_in']:.4f}in {page['right_in']:.4f}in {page['bottom_in']:.4f}in {page['left_in']:.4f}in;
  background: #ffffff;
  box-shadow: 0 0 8px rgba(0,0,0,0.25);
}}
.page p {{
  padding: 0;
}}
@media print {{
  html, body {{ background: #ffffff; }}
  .page {{ box-shadow: none; margin: 0; width: auto; min-height: 0; }}
}}
</style>
</head>
<body>
<div class="page">
{chr(10).join(body_parts)}
</div>
</body>
</html>
"""

with open(OUT, 'w', encoding='utf-8') as f:
    f.write(html_doc)

print("Wrote", OUT, "paragraphs:", len(body_parts))
